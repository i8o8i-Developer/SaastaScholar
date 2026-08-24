import os
import re
import html
from Apps.Academic.models import MaterialPage

def parse_markdown_to_html(raw_text, subject_code="", unit_number=1, page_number=1, total_pages=1, title=""):
    if not raw_text:
        raw_text = "Verified Curriculum Reference Material & Derivation Notes."

    lines = raw_text.split("\n")
    html_parts = []
    in_code_block = False
    code_lang = ""
    code_lines = []
    in_list = False

    for line in lines:
        stripped = line.strip()
        
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                code_content = html.escape("\n".join(code_lines))
                lang_label = f'<span class="text-[0.65rem] font-mono text-muted-foreground uppercase font-bold tracking-widest">{code_lang or "CODE"}</span>' if code_lang else ''
                html_parts.append(
                    f'<div class="my-4 rounded-2xl overflow-hidden border border-border/80 bg-zinc-950 shadow-md">'
                    f'<div class="flex items-center justify-between px-4 py-2 bg-zinc-900 border-b border-zinc-800">'
                    f'{lang_label}<span class="text-[0.65rem] font-mono text-emerald-400 font-bold">Copy Protected · DRM</span>'
                    f'</div><pre class="p-4 text-emerald-400 text-xs font-mono overflow-x-auto leading-relaxed"><code>{code_content}</code></pre></div>'
                )
                code_lines = []
                code_lang = ""
            else:
                if in_list:
                    html_parts.append("</ul>")
                    in_list = False
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Close list if not a list item
        if in_list and not (stripped.startswith("- ") or stripped.startswith("* ") or (len(stripped) > 2 and stripped[0].isdigit() and stripped[1] == ".")):
            html_parts.append("</ul>")
            in_list = False

        if not stripped:
            continue

        # Headers
        if stripped.startswith("# "):
            content = parse_inline(stripped[2:])
            html_parts.append(f'<h1 class="text-2xl sm:text-3xl font-black text-foreground mt-6 mb-3 pb-2 border-b border-border/80 tracking-tight">{content}</h1>')
        elif stripped.startswith("## "):
            content = parse_inline(stripped[3:])
            html_parts.append(f'<h2 class="text-xl font-bold text-foreground mt-5 mb-2.5 tracking-tight flex items-center gap-2.5"><span class="h-2.5 w-2.5 rounded-full bg-emerald-500 inline-block shrink-0"></span>{content}</h2>')
        elif stripped.startswith("### "):
            content = parse_inline(stripped[4:])
            html_parts.append(f'<h3 class="text-base font-bold text-emerald-600 dark:text-emerald-400 mt-4 mb-2 uppercase tracking-wider">{content}</h3>')
        elif stripped.startswith("#### "):
            content = parse_inline(stripped[5:])
            html_parts.append(f'<h4 class="text-sm font-bold text-foreground/90 mt-3 mb-1.5">{content}</h4>')
        elif stripped.startswith("> "):
            content = parse_inline(stripped[2:])
            html_parts.append(f'<blockquote class="my-3 pl-4 border-l-4 border-emerald-500/60 bg-emerald-500/5 py-2.5 pr-3 rounded-r-2xl text-xs sm:text-sm text-foreground/90 italic">{content}</blockquote>')
        elif stripped.startswith("- ") or stripped.startswith("* "):
            if not in_list:
                html_parts.append('<ul class="my-2.5 space-y-1.5">')
                in_list = True
            content = parse_inline(stripped[2:])
            html_parts.append(f'<li class="flex items-start gap-2 text-sm text-foreground/90 leading-relaxed"><span class="h-1.5 w-1.5 rounded-full bg-emerald-500/70 mt-2 shrink-0"></span><span>{content}</span></li>')
        elif len(stripped) > 2 and stripped[0].isdigit() and stripped[1] == ".":
            if not in_list:
                html_parts.append('<ul class="my-2.5 space-y-1.5">')
                in_list = True
            num = stripped.split(".")[0]
            rest = ".".join(stripped.split(".")[1:]).strip()
            content = parse_inline(rest)
            html_parts.append(f'<li class="flex items-start gap-2 text-sm text-foreground/90 leading-relaxed"><span class="font-mono text-xs font-bold text-emerald-600 dark:text-emerald-400 mt-0.5 shrink-0">{num}.</span><span>{content}</span></li>')
        elif stripped.startswith("---") or stripped.startswith("***"):
            html_parts.append('<hr class="my-6 border-border/80" />')
        else:
            content = parse_inline(stripped)
            html_parts.append(f'<p class="text-sm leading-relaxed text-foreground/90 my-2">{content}</p>')

    if in_list:
        html_parts.append("</ul>")

    if in_code_block:
        code_content = html.escape("\n".join(code_lines))
        html_parts.append(f'<pre class="p-4 bg-zinc-950 text-emerald-400 rounded-2xl text-xs font-mono overflow-x-auto my-3 border border-emerald-500/30"><code>{code_content}</code></pre>')

    body_html = "".join(html_parts)
    return f"""
    <div class="space-y-4 font-sans select-none">
      <div class="border-b border-border/70 pb-4">
        <div class="flex items-center gap-2 flex-wrap mb-2">
          <span class="px-2.5 py-0.5 rounded-full bg-emerald-500/10 border border-emerald-500/20 text-emerald-600 dark:text-emerald-400 font-mono text-[0.7rem] font-bold uppercase tracking-wider">
            {subject_code or "NIMS"}
          </span>
          <span class="px-2.5 py-0.5 rounded-full bg-surface border border-border text-muted-foreground font-mono text-[0.7rem] font-bold">
            UNIT {unit_number or 1}
          </span>
          <span class="px-2.5 py-0.5 rounded-full bg-surface border border-border text-muted-foreground font-mono text-[0.7rem] font-bold">
            PAGE {page_number} OF {total_pages}
          </span>
        </div>
        <h2 class="text-2xl sm:text-3xl font-black tracking-tight text-foreground">{title}</h2>
        <p class="text-xs text-muted-foreground mt-1 font-medium">Verified Academic Courseware · Real-time Forensic Telemetry Active</p>
      </div>
      <div class="space-y-2 pt-2">
        {body_html}
      </div>
    </div>
    """

def parse_inline(text):
    if not text:
        return ""
    # Bold **text**
    text = re.sub(r'\*\*(.*?)\*\*', r'<strong class="font-bold text-foreground">\1</strong>', text)
    # Inline Code `code`
    text = re.sub(r'`(.*?)`', r'<code class="px-1.5 py-0.5 rounded-md bg-emerald-500/10 text-emerald-700 dark:text-emerald-300 font-mono text-xs border border-emerald-500/20">\1</code>', text)
    # Italic *text*
    text = re.sub(r'(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)', r'<em class="italic text-foreground/85">\1</em>', text)
    # Math formulas like $O(N)$
    text = re.sub(r'\$(.*?)\$', r'<span class="font-mono text-emerald-600 dark:text-emerald-400 bg-emerald-500/10 px-1.5 py-0.5 rounded font-bold">\1</span>', text)
    return text

def extract_and_populate_pages(material):
    if not material.file:
        return

    try:
        file_path = material.file.path
        if not os.path.exists(file_path):
            return
        
        file_name = material.file.name.lower()
        subject_code = material.subject.code if material.subject else "NIMS"
        unit_num = material.unit_number or 1

        MaterialPage.objects.filter(material=material).delete()

        # Check if LibreOffice conversion is needed
        if file_name.endswith((".docx", ".doc", ".pptx", ".ppt")):
            try:
                import subprocess
                from django.core.files import File
                outdir = os.path.dirname(file_path)
                cmd = ["libreoffice", "--headless", "--convert-to", "pdf", file_path, "--outdir", outdir]
                subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                
                base_name = os.path.splitext(os.path.basename(file_path))[0]
                pdf_path = os.path.join(outdir, f"{base_name}.pdf")
                if os.path.exists(pdf_path):
                    with open(pdf_path, 'rb') as f:
                        material.file.save(f"{base_name}.pdf", File(f), save=True)
                    file_path = material.file.path
                    file_name = material.file.name.lower()
                    
                    # Optionally remove the intermediate PDF and original file from disk if desired
                    # For safety, we can just leave them as they are in MEDIA_ROOT
            except Exception as e:
                print("LibreOffice conversion failed:", e)

        # 1. PDF Handler
        if file_name.endswith(".pdf"):
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                total_pages = len(reader.pages)
                for idx, page in enumerate(reader.pages, start=1):
                    raw_text = page.extract_text() or ""
                    if not raw_text.strip():
                        raw_text = f"## Page {idx}: Visual Vector Manuscript\n\n- High-Yield Visual Diagram & Faculty Reference\n- Digital Copyright Protected by NimsWallah DRM Pipeline"
                    
                    formatted_html = parse_markdown_to_html(
                        raw_text=raw_text,
                        subject_code=subject_code,
                        unit_number=unit_num,
                        page_number=idx,
                        total_pages=total_pages,
                        title=f"{material.title} (Page {idx})",
                    )
                    MaterialPage.objects.create(
                        material=material,
                        page_number=idx,
                        page_title=f"Page {idx}: {material.title}",
                        content_html=formatted_html,
                    )
                if total_pages > 0:
                    material.page_count = total_pages
                    material.save(update_fields=["page_count"])
                return
            except Exception as e:
                print("PDF Extraction Error:", e)

        # 2. PPTX / PPT Presentation Handler
        elif file_name.endswith((".pptx", ".ppt")):
            try:
                import pptx
                prs = pptx.Presentation(file_path)
                total_slides = len(prs.slides)
                for idx, slide in enumerate(prs.slides, start=1):
                    slide_title = f"Slide {idx}"
                    if slide.shapes.title and slide.shapes.title.text.strip():
                        slide_title = slide.shapes.title.text.strip()
                    
                    bullet_points = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                txt = paragraph.text.strip()
                                if txt and txt != slide_title:
                                    bullet_points.append(f"- {txt}")
                        elif shape.has_table:
                            for row in shape.table.rows:
                                row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                                if row_txt:
                                    bullet_points.append(f"- **Table Row**: {row_txt}")

                    if not bullet_points:
                        bullet_points = ["- High-Yield Presentation Slide Graphic & Key Concepts", "- Faculty Verified Lecture Reference Deck"]

                    raw_md = f"## {slide_title}\n\n" + "\n".join(bullet_points)
                    formatted_html = parse_markdown_to_html(
                        raw_text=raw_md,
                        subject_code=subject_code,
                        unit_number=unit_num,
                        page_number=idx,
                        total_pages=total_slides,
                        title=f"{material.title} - Slide {idx}",
                    )
                    MaterialPage.objects.create(
                        material=material,
                        page_number=idx,
                        page_title=f"Slide {idx}: {slide_title}",
                        content_html=formatted_html,
                    )
                if total_slides > 0:
                    material.page_count = total_slides
                    material.save(update_fields=["page_count"])
                return
            except Exception as e:
                print("PPTX Extraction Error:", e)

        # 3. DOCX Word Handler
        elif file_name.endswith((".docx", ".doc")):
            try:
                import docx
                doc = docx.Document(file_path)
                paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                
                # Also extract table text
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if row_cells:
                            paras.append(f"- **Data Row**: {' | '.join(row_cells)}")

                raw_md = "\n\n".join(paras) if paras else (material.summary or "Verified Academic Document.")
                formatted_html = parse_markdown_to_html(
                    raw_text=raw_md,
                    subject_code=subject_code,
                    unit_number=unit_num,
                    page_number=1,
                    total_pages=1,
                    title=material.title,
                )
                MaterialPage.objects.create(
                    material=material,
                    page_number=1,
                    page_title=f"Page 1: {material.title}",
                    content_html=formatted_html,
                )
                material.page_count = 1
                material.save(update_fields=["page_count"])
                return
            except Exception as e:
                print("DOCX Extraction Error:", e)

        # 4. MD & TXT Handler
        elif file_name.endswith((".txt", ".md")):
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    full_txt = f.read()
                
                formatted_html = parse_markdown_to_html(
                    raw_text=full_txt or material.summary or "Verified Document Content.",
                    subject_code=subject_code,
                    unit_number=unit_num,
                    page_number=1,
                    total_pages=1,
                    title=material.title,
                )
                MaterialPage.objects.create(
                    material=material,
                    page_number=1,
                    page_title=f"Page 1: {material.title}",
                    content_html=formatted_html,
                )
                material.page_count = 1
                material.save(update_fields=["page_count"])
                return
            except Exception as e:
                print("TXT/MD Extraction Error:", e)

        # Fallback default
        formatted_html = parse_markdown_to_html(
            raw_text=material.summary or "Verified Courseware Document.",
            subject_code=subject_code,
            unit_number=unit_num,
            page_number=1,
            total_pages=1,
            title=material.title,
        )
        MaterialPage.objects.create(
            material=material,
            page_number=1,
            page_title=f"Page 1: {material.title}",
            content_html=formatted_html,
        )
    except Exception as e:
        print("General Document Extraction Error :", e)
